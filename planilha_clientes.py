"""
Integração com planilha de clientes
====================================

Lê uma planilha .xlsx com uma linha por rescisão (uma coluna por campo de
entrada de `DadosRescisao`), calcula cada uma com
`calculadora_rescisao.calcular_rescisao()` e grava uma planilha de saída com
os totais e alertas de cada linha.

Ajuste `COLUNAS_ENTRADA` abaixo ao layout real da sua planilha de clientes —
os nomes de coluna aqui são um ponto de partida, não um padrão fixo. Campos
mais avançados de `DadosRescisao` (estabilidades, contrato de experiência,
horas extras detalhadas etc.) não têm coluna própria aqui; adicione conforme
a necessidade.

Uso:
    python planilha_clientes.py modelo.xlsx              # gera planilha modelo
    python planilha_clientes.py entrada.xlsx saida.xlsx   # processa e calcula
"""

import sys
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from calculadora_rescisao import DadosRescisao, calcular_rescisao

COLUNAS_ENTRADA = [
    "nome",
    "cpf",
    "salario_base",
    "data_admissao",
    "data_desligamento",
    "tipo_rescisao",
    "aviso_previo",
    "ferias_vencidas",
    "dependentes_irrf",
    "saldo_fgts_depositado",
    "media_horas_extras_mensal",
    "numero_solicitacoes_seguro_desemprego_anteriores",
]

COLUNAS_SAIDA = [
    "total_proventos",
    "total_descontos",
    "valor_liquido",
    "prazo_pagamento",
    "seguro_desemprego_parcelas",
    "seguro_desemprego_valor_total",
    "alertas_risco",
    "erro",
]

TIPOS_RESCISAO_VALIDOS = {
    "sem_justa_causa",
    "pedido_demissao",
    "justa_causa",
    "acordo_484a",
    "termino_experiencia",
    "rescisao_indireta",
}


def gerar_planilha_modelo(caminho: str) -> None:
    """Cria uma planilha modelo com o cabeçalho esperado e uma linha de exemplo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rescisões"
    ws.append(COLUNAS_ENTRADA)
    ws.append([
        "Fulano de Tal", "000.000.000-00", 3500.00, date(2022, 3, 10), date(2026, 7, 15),
        "sem_justa_causa", "indenizado", "não", 0, 6200.00, 0, 0,
    ])
    wb.save(caminho)


def _parse_data(valor) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return datetime.strptime(str(valor).strip(), "%d/%m/%Y").date()


def _parse_bool(valor) -> bool:
    return str(valor).strip().lower() in ("sim", "true", "1", "verdadeiro")


def _linha_para_dados(linha: dict) -> DadosRescisao:
    tipo = str(linha["tipo_rescisao"]).strip()
    if tipo not in TIPOS_RESCISAO_VALIDOS:
        raise ValueError(f"tipo_rescisao inválido: {tipo!r}")
    return DadosRescisao(
        salario_base=float(linha["salario_base"]),
        data_admissao=_parse_data(linha["data_admissao"]),
        data_desligamento=_parse_data(linha["data_desligamento"]),
        tipo_rescisao=tipo,
        aviso_previo=str(linha.get("aviso_previo") or "indenizado").strip(),
        ferias_vencidas=_parse_bool(linha.get("ferias_vencidas", "não")),
        dependentes_irrf=int(linha.get("dependentes_irrf") or 0),
        saldo_fgts_depositado=float(linha.get("saldo_fgts_depositado") or 0.0),
        media_horas_extras_mensal=float(linha.get("media_horas_extras_mensal") or 0.0),
        numero_solicitacoes_seguro_desemprego_anteriores=int(
            linha.get("numero_solicitacoes_seguro_desemprego_anteriores") or 0
        ),
    )


def processar_planilha(caminho_entrada: str, caminho_saida: str) -> None:
    """Lê `caminho_entrada`, calcula cada rescisão e grava `caminho_saida` com
    os totais ao lado dos dados originais. Linhas com erro (dado inválido,
    tipo de rescisão desconhecido etc.) são mantidas com a coluna `erro`
    preenchida, em vez de interromper o processamento das demais.
    """
    wb_in = load_workbook(caminho_entrada)
    ws_in = wb_in.active
    cabecalho = [c.value for c in next(ws_in.iter_rows(min_row=1, max_row=1))]

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "Resultado"
    ws_out.append(cabecalho + COLUNAS_SAIDA)

    for linha_valores in ws_in.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in linha_valores):
            continue
        linha = dict(zip(cabecalho, linha_valores))
        try:
            dados = _linha_para_dados(linha)
            resultado = calcular_rescisao(dados)
            saida = [
                resultado["total_proventos"],
                resultado["total_descontos"],
                resultado["valor_liquido"],
                resultado["prazo_pagamento"].strftime("%d/%m/%Y"),
                resultado["seguro_desemprego"]["numero_parcelas"],
                resultado["seguro_desemprego"]["valor_total_estimado"],
                "; ".join(resultado["alertas_risco"]),
                "",
            ]
        except Exception as exc:
            saida = ["", "", "", "", "", "", "", str(exc)]
        ws_out.append(list(linha_valores) + saida)

    wb_out.save(caminho_saida)


if __name__ == "__main__":
    if len(sys.argv) == 2:
        gerar_planilha_modelo(sys.argv[1])
        print(f"Planilha modelo criada em {sys.argv[1]}")
    elif len(sys.argv) == 3:
        processar_planilha(sys.argv[1], sys.argv[2])
        print(f"Planilha processada, resultado em {sys.argv[2]}")
    else:
        print("Uso:")
        print("  python planilha_clientes.py modelo.xlsx")
        print("  python planilha_clientes.py entrada.xlsx saida.xlsx")
