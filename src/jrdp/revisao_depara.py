"""Planilha de revisão humana da camada de identidade.

Gera localmente (nunca no Git) uma planilha `.xlsx` com uma linha por
pessoa/unidade não encontrada, com sugestão fuzzy quando existir, para um
analista de DP revisar. Nenhuma linha vira de-para automaticamente — só
as marcadas explicitamente "APROVAR" pelo analista, com aprovador e data
preenchidos, são importadas.

Fluxo (ver docs/DECISIONS.md, 2026-09-17):

    60 não resolvidos (identidade.coletar_ocorrencias_nao_encontradas)
        -> montar_linhas_revisao (sugestão fuzzy só como ponto de partida)
        -> escrever_planilha_revisao (.xlsx local, fora do Git)
        -> [revisão humana preenche "Decisão analista"/"Aprovado por"/
            "Data aprovação" na planilha]
        -> importar_decisoes_aprovadas (só linhas "APROVAR")
        -> depara.mesclar_depara + depara.salvar_depara
        -> reexecutar o pipeline
"""

from dataclasses import dataclass

from .cadastro_ativos import RegistroCadastro
from .depara import DeParaContractError, RegistroDePara
from .sugestao_fuzzy import sugerir_candidatos

COLUNAS = [
    "Nome origem",
    "Unidade",
    "Eventos",
    "Sugestão",
    "Código sugerido",
    "Nome cadastro",
    "Confiança diagnóstica",
    "Decisão analista",
    "Aprovado por",
    "Data aprovação",
    "Observação",
]

DECISAO_APROVAR = "APROVAR"
DECISAO_REJEITAR = "REJEITAR"


@dataclass(frozen=True)
class LinhaRevisao:
    nome_origem: str
    unidade: str
    eventos: list[int]
    sugestao: str  # "Automática" ou "Manual"
    codigo_sugerido: str | None
    nome_cadastro: str | None
    confianca: float | None


def montar_linhas_revisao(
    ocorrencias: list[dict], cadastro: list[RegistroCadastro]
) -> list[LinhaRevisao]:
    """`ocorrencias` vem de
    `identidade.coletar_ocorrencias_nao_encontradas`. Para cada uma,
    tenta uma sugestão fuzzy (só diagnóstico) e monta a linha da planilha.
    """
    linhas = []
    for ocorrencia in ocorrencias:
        sugestoes = sugerir_candidatos(ocorrencia["nome_origem"], cadastro, limite=1)
        if sugestoes:
            candidato = sugestoes[0]
            linhas.append(
                LinhaRevisao(
                    nome_origem=ocorrencia["nome_origem"],
                    unidade=ocorrencia["unidade"],
                    eventos=ocorrencia["eventos"],
                    sugestao="Automática",
                    codigo_sugerido=candidato.contrato,
                    nome_cadastro=candidato.nome_cadastro,
                    confianca=candidato.similaridade,
                )
            )
        else:
            linhas.append(
                LinhaRevisao(
                    nome_origem=ocorrencia["nome_origem"],
                    unidade=ocorrencia["unidade"],
                    eventos=ocorrencia["eventos"],
                    sugestao="Manual",
                    codigo_sugerido=None,
                    nome_cadastro=None,
                    confianca=None,
                )
            )
    return linhas


def escrever_planilha_revisao(linhas: list[LinhaRevisao], caminho) -> None:
    """Escreve a planilha de revisão em `.xlsx`. Chamar sempre com um
    caminho fora do Git (ex.: dentro de `homologacao/`) — a planilha
    carrega nomes reais."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revisão de identidade"
    ws.append(COLUNAS)
    for linha in linhas:
        ws.append(
            [
                linha.nome_origem,
                linha.unidade,
                ", ".join(str(e) for e in linha.eventos),
                linha.sugestao,
                linha.codigo_sugerido or "",
                linha.nome_cadastro or "",
                f"{linha.confianca:.2f}" if linha.confianca is not None else "",
                "",  # Decisão analista — o analista preenche
                "",  # Aprovado por — o analista preenche
                "",  # Data aprovação — o analista preenche
                "",  # Observação — o analista preenche
            ]
        )
    wb.save(caminho)


def importar_decisoes_aprovadas(caminho, cliente: str) -> list[RegistroDePara]:
    """Lê a planilha já revisada e devolve só as entradas de-para das
    linhas marcadas `"APROVAR"`. Nunca importa automaticamente uma linha
    sem decisão explícita, sem aprovador ou sem data — isso é erro de
    preenchimento, não algo para assumir.
    """
    import openpyxl

    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb.active
    cabecalho = [celula.value for celula in next(ws.iter_rows(min_row=1, max_row=1))]
    if cabecalho != COLUNAS:
        raise DeParaContractError(
            f"Cabeçalho da planilha de revisão não bate com o esperado. "
            f"Esperado {COLUNAS!r}, recebido {cabecalho!r}."
        )
    indice = {nome: i for i, nome in enumerate(cabecalho)}

    registros = []
    for numero_linha, linha in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        decisao = str(linha[indice["Decisão analista"]] or "").strip().upper()
        if decisao != DECISAO_APROVAR:
            continue

        codigo = linha[indice["Código sugerido"]]
        nome_cadastro = linha[indice["Nome cadastro"]]
        aprovado_por = linha[indice["Aprovado por"]]
        aprovado_em = linha[indice["Data aprovação"]]
        if not codigo or not aprovado_por or not aprovado_em:
            raise DeParaContractError(
                f"Linha {numero_linha} marcada 'APROVAR' sem código/aprovador/"
                "data preenchidos — corrija a planilha antes de importar."
            )

        registros.append(
            RegistroDePara(
                nome_origem=linha[indice["Nome origem"]],
                unidade=linha[indice["Unidade"]],
                cliente=cliente,
                codigo_questor=str(codigo),
                nome_canonico=str(nome_cadastro) if nome_cadastro else linha[indice["Nome origem"]],
                status="aprovado",
                evidencia=str(linha[indice["Observação"]] or "Aprovado via planilha de revisão de identidade"),
                aprovado_por=str(aprovado_por),
                aprovado_em=str(aprovado_em),
            )
        )
    return registros
