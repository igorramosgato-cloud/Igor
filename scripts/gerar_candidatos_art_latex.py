"""Gera os CANDIDATOS de homologação tipo V da ART LATEX (nunca produção).

Substitui os scripts ad-hoc usados manualmente durante a homologação do
P01 (ver docs/DECISIONS.md, 2026-09-18 em diante). Roda contra os
arquivos reais em homologacao/art_latex/questor/origem/ (nunca versionados
no Git) e o de-para já homologado em
homologacao/art_latex/questor/depara/depara_nomes.json.

Uso:
    PYTHONPATH=src python3 scripts/gerar_candidatos_art_latex.py

Nunca gera arquivo de produção — só candidatos de homologação, nomeados
explicitamente, salvos fora do Git em
homologacao/art_latex/questor/homologacao_v_candidatos/. A importação real
no Questor continua sendo manual (ver docs/P01_CHECKLIST_IMPORTACAO_MANUAL_V.md).
"""

import hashlib
import json
import os
import sys
from decimal import Decimal

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

import openpyxl

from jrdp.cadastro_ativos import parse_base_ativos
from jrdp.config import load_cliente_config
from jrdp.decimais import valor_origem_para_decimal
from jrdp.depara import carregar_depara
from jrdp.exportadores import ExportacaoBlockedError, QuestorExporterColunar
from jrdp.extratores.cesta_basica import extrair_nomes_cesta_basica
from jrdp.extratores.valor_simples import extrair_valor_simples
from jrdp.conferencia import gerar_relatorio_evento
from jrdp.pipeline import construir_lancamentos_cesta_basica, construir_lancamentos_evento
from jrdp.questor_layout import parse_arquivo_layout_colunar

CLIENTE = "ART LATEX"
COMPETENCIA = "08/2026"

BASE = os.path.join(
    os.path.dirname(__file__), "..", "homologacao", "art_latex", "questor"
)
ARQ_MATRIZ = os.path.join(BASE, "origem", "planilha_importacao_matriz.xlsm")
ARQ_FILIAL = os.path.join(BASE, "origem", "planilha_importacao_filial.xlsm")
ARQ_CADASTRO = os.path.join(BASE, "origem", "base_ativos_art_latex.csv")
DEPARA_PATH = os.path.join(BASE, "depara", "depara_nomes.json")
CANDIDATOS_DIR = os.path.join(BASE, "homologacao_v_candidatos")

# (codigo_evento, aba_origem, rotulo) — eventos de valor com código
# confirmado para Matriz+Filial (ver docs/DECISIONS.md, 2026-09-18). Ordem
# igual ao checklist docs/P01_CHECKLIST_IMPORTACAO_MANUAL_V.md: 806, 813,
# 1524 (Cesta, extração diferente), 1955.
EVENTOS_VALOR_SIMPLES_ANTES_CESTA = [
    (806, "convenio farmacia", "Desconto Farmácia"),
    (813, "Vale-compras", "Desconto Compras"),
]
EVENTO_CESTA = (1524, "Desconto Cesta Basica")
EVENTOS_VALOR_SIMPLES_APOS_CESTA = [
    (1955, "Vale-refeicao", "Vale Refeição"),
]


def linhas_aba(wb, aba, min_row=6):
    ws = wb[aba]
    linhas = []
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        if row[1] is None and (len(row) <= 2 or row[2] is None):
            continue
        linhas.append(row)
    return linhas


def main():
    os.makedirs(CANDIDATOS_DIR, exist_ok=True)

    with open(ARQ_CADASTRO, "rb") as f:
        cadastro = parse_base_ativos(f.read())
    config_cliente = load_cliente_config("art_latex")
    depara = carregar_depara(DEPARA_PATH)
    print(f"De-para carregado: {len(depara)} entradas homologadas")

    wb_matriz = openpyxl.load_workbook(ARQ_MATRIZ, read_only=True, data_only=True)
    wb_filial = openpyxl.load_workbook(ARQ_FILIAL, read_only=True, data_only=True)

    entradas_exportador = []
    relatorios = {}

    def _adicionar_valor_simples(codigo_evento, aba, rotulo):
        registros = extrair_valor_simples(
            linhas_aba(wb_filial, aba), unidade="filial",
            arquivo_origem="planilha_importacao_filial.xlsm", aba_origem=aba,
        ) + extrair_valor_simples(
            linhas_aba(wb_matriz, aba), unidade="matriz",
            arquivo_origem="planilha_importacao_matriz.xlsm", aba_origem=aba,
        )
        lancamentos = construir_lancamentos_evento(
            registros, cadastro, config_cliente, CLIENTE, COMPETENCIA, codigo_evento, depara=depara
        )
        relatorio = gerar_relatorio_evento(lancamentos)
        relatorios[codigo_evento] = relatorio
        entradas_exportador.append((str(codigo_evento), "V", rotulo, lancamentos, relatorio))

    for codigo_evento, aba, rotulo in EVENTOS_VALOR_SIMPLES_ANTES_CESTA:
        _adicionar_valor_simples(codigo_evento, aba, rotulo)

    codigo_cesta, rotulo_cesta = EVENTO_CESTA
    nomes_matriz = extrair_nomes_cesta_basica(linhas_aba(wb_matriz, "Cesta basica"))
    nomes_filial = extrair_nomes_cesta_basica(linhas_aba(wb_filial, "Cesta basica"))
    lancamentos_cesta = construir_lancamentos_cesta_basica(
        nomes_matriz, nomes_filial, cadastro, config_cliente, CLIENTE, COMPETENCIA, codigo_cesta,
        arquivo_matriz="planilha_importacao_matriz.xlsm",
        arquivo_filial="planilha_importacao_filial.xlsm",
        depara=depara,
    )
    relatorio_cesta = gerar_relatorio_evento(lancamentos_cesta)
    relatorios[codigo_cesta] = relatorio_cesta
    entradas_exportador.append((str(codigo_cesta), "V", rotulo_cesta, lancamentos_cesta, relatorio_cesta))

    for codigo_evento, aba, rotulo in EVENTOS_VALOR_SIMPLES_APOS_CESTA:
        _adicionar_valor_simples(codigo_evento, aba, rotulo)

    print("\nRelatório de conferência (identidade/matching) por evento:")
    for codigo_evento in sorted(relatorios):
        r = relatorios[codigo_evento]
        print(
            f"  {codigo_evento}: total={r.quantidade_total} "
            f"nao_encontrados={r.nao_encontrados} ambiguos={r.ambiguos} "
            f"status={r.status}"
        )

    exporter = QuestorExporterColunar()
    try:
        conteudo = exporter.exportar(entradas_exportador)
    except ExportacaoBlockedError as exc:
        print(f"\nNenhum candidato gerado: {exc}")
        return

    eventos_incluidos = [
        codigo for codigo, _tipo, _rotulo, _lancs, rel in entradas_exportador if rel.pode_exportar()
    ]
    eventos_excluidos = [
        codigo for codigo, _tipo, _rotulo, _lancs, rel in entradas_exportador if not rel.pode_exportar()
    ]

    nome_arquivo = (
        f"CANDIDATO_HOMOLOGACAO_COLUNAR_{'_'.join(eventos_incluidos)}"
        f"_competencia_{COMPETENCIA.replace('/', '-')}.csv"
    )
    caminho = os.path.join(CANDIDATOS_DIR, nome_arquivo)
    with open(caminho, "wb") as f:
        f.write(conteudo)
    sha = hashlib.sha256(conteudo).hexdigest()

    # validação round-trip + reconciliação (quantidade e valor, via Decimal)
    arquivo_parseado = parse_arquivo_layout_colunar(conteudo)
    reconciliacao = {}
    for codigo in eventos_incluidos:
        qtd = sum(1 for r in arquivo_parseado.registros if codigo in r.valores)
        total = sum(
            (valor_origem_para_decimal(r.valores[codigo].replace(",", ".")) for r in arquivo_parseado.registros if codigo in r.valores),
            start=Decimal("0"),
        )
        relatorio = relatorios[int(codigo)]
        assert qtd == relatorio.quantidade_total, f"Divergência de quantidade no evento {codigo}"
        assert total == relatorio.soma_valor, f"Divergência de valor no evento {codigo}"
        reconciliacao[codigo] = {"quantidade": qtd, "total": str(total)}

    print(f"\nCandidato gerado: {nome_arquivo}")
    print(f"SHA-256: {sha}")
    print(f"Colaboradores (linhas): {len(arquivo_parseado.registros)}")
    print(f"Eventos incluídos: {eventos_incluidos}")
    print(f"Eventos excluídos (BLOCKED): {eventos_excluidos}")
    print("Reconciliação (quantidade e valor batem com o relatório de conferência):")
    for codigo, dado in reconciliacao.items():
        print(f"  {codigo}: {dado}")

    manifesto = {
        "cliente": CLIENTE,
        "competencia": COMPETENCIA,
        "natureza": "CANDIDATO DE HOMOLOGACAO — NAO E PRODUCAO, NAO IMPORTAR AUTOMATICAMENTE",
        "arquivo": nome_arquivo,
        "sha256": sha,
        "eventos_incluidos": eventos_incluidos,
        "eventos_excluidos": eventos_excluidos,
        "reconciliacao": reconciliacao,
    }
    caminho_manifesto = os.path.join(CANDIDATOS_DIR, "manifesto_candidato_colunar.json")
    with open(caminho_manifesto, "w", encoding="utf-8") as f:
        json.dump(manifesto, f, ensure_ascii=False, indent=2)
    print(f"\nManifesto salvo em {caminho_manifesto}")


if __name__ == "__main__":
    main()
