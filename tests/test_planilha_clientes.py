from openpyxl import load_workbook

from planilha_clientes import gerar_planilha_modelo, processar_planilha


def test_gerar_modelo_e_processar_planilha(tmp_path):
    modelo = tmp_path / "modelo.xlsx"
    saida = tmp_path / "saida.xlsx"

    gerar_planilha_modelo(str(modelo))
    assert modelo.exists()

    processar_planilha(str(modelo), str(saida))
    assert saida.exists()

    wb = load_workbook(saida)
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    cabecalho = linhas[0]
    dados = dict(zip(cabecalho, linhas[1]))

    assert dados["erro"] in (None, "")
    assert dados["valor_liquido"] > 0
    assert dados["total_proventos"] > dados["total_descontos"]


def test_linha_invalida_gera_erro_sem_derrubar_processamento(tmp_path):
    from openpyxl import Workbook

    entrada = tmp_path / "entrada.xlsx"
    saida = tmp_path / "saida.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.append([
        "nome", "cpf", "salario_base", "data_admissao", "data_desligamento",
        "tipo_rescisao", "aviso_previo", "ferias_vencidas", "dependentes_irrf",
        "saldo_fgts_depositado", "media_horas_extras_mensal",
        "numero_solicitacoes_seguro_desemprego_anteriores",
    ])
    ws.append([
        "Linha Ruim", "111", 3000, "10/03/2022", "15/07/2026",
        "tipo_invalido", "indenizado", "não", 0, 0, 0, 0,
    ])
    wb.save(entrada)

    processar_planilha(str(entrada), str(saida))

    wb_out = load_workbook(saida)
    ws_out = wb_out.active
    linhas = list(ws_out.iter_rows(values_only=True))
    cabecalho = linhas[0]
    dados = dict(zip(cabecalho, linhas[1]))
    assert dados["erro"]
