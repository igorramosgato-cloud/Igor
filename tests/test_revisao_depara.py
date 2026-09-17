"""Testes da planilha de revisão de identidade. Dados 100% fictícios,
arquivos escritos em tmp_path (nunca em homologacao/ real) — ver
.claude/rules/homologacao-dados.md.
"""

import pytest

from jrdp.cadastro_ativos import RegistroCadastro
from jrdp.depara import DeParaContractError
from jrdp.revisao_depara import (
    COLUNAS,
    escrever_planilha_revisao,
    importar_decisoes_aprovadas,
    montar_linhas_revisao,
)

CLIENTE = "CLIENTE_FICTICIO"


def _cadastro(contrato, nome):
    return RegistroCadastro(
        contrato=contrato, nome=nome, admissao="01/01/2020", descricao="CARGO", cpf="0"
    )


CADASTRO = [
    _cadastro("1", "JOAO DA SILVA SANTOS"),
    _cadastro("2", "MARIA APARECIDA DE SOUZA"),
]


def _ocorrencia(nome, unidade="filial", eventos=(1955,)):
    return {"nome_origem": nome, "unidade": unidade, "eventos": list(eventos)}


def test_monta_linha_com_sugestao_automatica():
    linhas = montar_linhas_revisao([_ocorrencia("JOAO DA SILVA")], CADASTRO)
    assert len(linhas) == 1
    assert linhas[0].sugestao == "Automática"
    assert linhas[0].codigo_sugerido == "1"
    assert linhas[0].nome_cadastro == "JOAO DA SILVA SANTOS"


def test_monta_linha_manual_quando_sem_candidato():
    linhas = montar_linhas_revisao([_ocorrencia("XYZXYZXYZ NAO PARECIDO")], CADASTRO)
    assert linhas[0].sugestao == "Manual"
    assert linhas[0].codigo_sugerido is None


def test_eventos_agregados_na_linha():
    linhas = montar_linhas_revisao([_ocorrencia("JOAO DA SILVA", eventos=(1955, 813))], CADASTRO)
    assert linhas[0].eventos == [1955, 813]


def test_escreve_e_le_cabecalho_da_planilha(tmp_path):
    linhas = montar_linhas_revisao([_ocorrencia("JOAO DA SILVA")], CADASTRO)
    caminho = tmp_path / "revisao.xlsx"
    escrever_planilha_revisao(linhas, caminho)

    import openpyxl

    wb = openpyxl.load_workbook(caminho)
    ws = wb.active
    cabecalho = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert cabecalho == COLUNAS


def _preencher_decisao(caminho, linha_numero, decisao, aprovado_por=None, aprovado_em=None, observacao=None):
    import openpyxl

    wb = openpyxl.load_workbook(caminho)
    ws = wb.active
    idx = {c.value: i + 1 for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}
    ws.cell(row=linha_numero, column=idx["Decisão analista"], value=decisao)
    if aprovado_por is not None:
        ws.cell(row=linha_numero, column=idx["Aprovado por"], value=aprovado_por)
    if aprovado_em is not None:
        ws.cell(row=linha_numero, column=idx["Data aprovação"], value=aprovado_em)
    if observacao is not None:
        ws.cell(row=linha_numero, column=idx["Observação"], value=observacao)
    wb.save(caminho)


def test_importa_apenas_linhas_aprovadas(tmp_path):
    ocorrencias = [_ocorrencia("JOAO DA SILVA"), _ocorrencia("MARIA APARECIDA")]
    linhas = montar_linhas_revisao(ocorrencias, CADASTRO)
    caminho = tmp_path / "revisao.xlsx"
    escrever_planilha_revisao(linhas, caminho)

    _preencher_decisao(caminho, 2, "APROVAR", "ANALISTA X", "2026-09-17", "Confirmado por CPF")
    # linha 3 (MARIA APARECIDA) fica sem decisão -- não deve ser importada

    registros = importar_decisoes_aprovadas(caminho, CLIENTE)
    assert len(registros) == 1
    assert registros[0].nome_origem == "JOAO DA SILVA"
    assert registros[0].codigo_questor == "1"
    assert registros[0].cliente == CLIENTE
    assert registros[0].aprovado_por == "ANALISTA X"


def test_linha_rejeitada_nao_e_importada(tmp_path):
    linhas = montar_linhas_revisao([_ocorrencia("JOAO DA SILVA")], CADASTRO)
    caminho = tmp_path / "revisao.xlsx"
    escrever_planilha_revisao(linhas, caminho)
    _preencher_decisao(caminho, 2, "REJEITAR")

    registros = importar_decisoes_aprovadas(caminho, CLIENTE)
    assert registros == []


def test_linha_sem_decisao_nao_e_importada(tmp_path):
    linhas = montar_linhas_revisao([_ocorrencia("JOAO DA SILVA")], CADASTRO)
    caminho = tmp_path / "revisao.xlsx"
    escrever_planilha_revisao(linhas, caminho)
    # nenhuma decisão preenchida

    registros = importar_decisoes_aprovadas(caminho, CLIENTE)
    assert registros == []


def test_aprovar_sem_aprovador_e_erro(tmp_path):
    linhas = montar_linhas_revisao([_ocorrencia("JOAO DA SILVA")], CADASTRO)
    caminho = tmp_path / "revisao.xlsx"
    escrever_planilha_revisao(linhas, caminho)
    _preencher_decisao(caminho, 2, "APROVAR")  # sem aprovado_por/aprovado_em

    with pytest.raises(DeParaContractError):
        importar_decisoes_aprovadas(caminho, CLIENTE)


def test_aprovar_caso_manual_com_correcao_do_analista(tmp_path):
    """Quando não há sugestão automática, o analista pode preencher
    'Código sugerido'/'Nome cadastro' manualmente antes de aprovar."""
    linhas = montar_linhas_revisao([_ocorrencia("XYZXYZXYZ NAO PARECIDO")], CADASTRO)
    caminho = tmp_path / "revisao.xlsx"
    escrever_planilha_revisao(linhas, caminho)

    import openpyxl

    wb = openpyxl.load_workbook(caminho)
    ws = wb.active
    idx = {c.value: i + 1 for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}
    ws.cell(row=2, column=idx["Código sugerido"], value="2")
    ws.cell(row=2, column=idx["Nome cadastro"], value="MARIA APARECIDA DE SOUZA")
    wb.save(caminho)
    _preencher_decisao(caminho, 2, "APROVAR", "ANALISTA X", "2026-09-17", "Encontrado manualmente")

    registros = importar_decisoes_aprovadas(caminho, CLIENTE)
    assert len(registros) == 1
    assert registros[0].codigo_questor == "2"


def test_cabecalho_alterado_e_rejeitado(tmp_path):
    linhas = montar_linhas_revisao([_ocorrencia("JOAO DA SILVA")], CADASTRO)
    caminho = tmp_path / "revisao.xlsx"
    escrever_planilha_revisao(linhas, caminho)

    import openpyxl

    wb = openpyxl.load_workbook(caminho)
    ws = wb.active
    ws.cell(row=1, column=1, value="Coluna renomeada")
    wb.save(caminho)

    with pytest.raises(DeParaContractError):
        importar_decisoes_aprovadas(caminho, CLIENTE)
